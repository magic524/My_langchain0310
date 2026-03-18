from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.*)$")
_SEPARATOR_CHARS = {"|", ":", "-", " "}


def _extract_tables_with_context(markdown_text: str) -> list[tuple[str, list[str], list[list[str]]]]:
    lines = markdown_text.splitlines()
    tables: list[tuple[str, list[str], list[list[str]]]] = []
    current_section = "Table"
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        heading_match = _HEADING_PATTERN.match(stripped)
        if heading_match:
            current_section = heading_match.group(2).strip()
            index += 1
            continue
        if _is_table_start(lines, index):
            headers, rows, consumed = _parse_table(lines, index)
            tables.append((current_section, headers, rows))
            index += consumed
            continue
        index += 1
    return tables


def _is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    first = lines[index].strip()
    second = lines[index + 1].strip()
    if "|" not in first or "|" not in second:
        return False
    return _is_separator_row(second)


def _is_separator_row(line: str) -> bool:
    stripped = line.strip().strip("|")
    return bool(stripped) and all(char in _SEPARATOR_CHARS for char in stripped)


def _parse_table(lines: list[str], start_index: int) -> tuple[list[str], list[list[str]], int]:
    raw_rows: list[str] = []
    index = start_index
    while index < len(lines):
        candidate = lines[index].strip()
        if not candidate.startswith("|"):
            break
        raw_rows.append(candidate)
        index += 1

    headers = _split_markdown_row(raw_rows[0])
    body_rows = [_split_markdown_row(row) for row in raw_rows[2:]]
    normalized_rows = [_normalize_row(row, len(headers)) for row in body_rows]
    return headers, normalized_rows, index - start_index


def _split_markdown_row(row: str) -> list[str]:
    stripped = row.strip().strip("|")
    cells: list[str] = []
    current: list[str] = []
    escaped = False
    for char in stripped:
        if escaped:
            current.append(char)
            escaped = False
            continue
        if char == "\\":
            escaped = True
            current.append(char)
            continue
        if char == "|":
            cells.append("".join(current).strip())
            current = []
            continue
        current.append(char)
    cells.append("".join(current).strip())
    return [cell.replace("<br>", "\n").replace("<br/>", "\n") for cell in cells]


def _normalize_row(row: list[str], width: int) -> list[str]:
    if len(row) < width:
        return row + [""] * (width - len(row))
    if len(row) > width:
        return row[: width - 1] + [" | ".join(row[width - 1 :])]
    return row


def export_markdown_to_xlsx(markdown_path: Path, xlsx_path: Path) -> None:
    markdown_text = markdown_path.read_text(encoding="utf-8")
    tables = _extract_tables_with_context(markdown_text)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "report_tables"
    ws.freeze_panes = "A2"

    title_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    header_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    title_font = Font(bold=True, size=12)
    header_font = Font(bold=True)
    wrap_alignment = Alignment(vertical="top", wrap_text=True)

    row_cursor = 1
    max_col_count = 1
    for table_index, (section_title, headers, rows) in enumerate(tables, start=1):
        table_title = f"{table_index}. {section_title}"
        ws.cell(row=row_cursor, column=1, value=table_title)
        ws.cell(row=row_cursor, column=1).font = title_font
        ws.cell(row=row_cursor, column=1).fill = title_fill
        ws.cell(row=row_cursor, column=1).alignment = wrap_alignment
        row_cursor += 1

        for col_index, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_cursor, column=col_index, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap_alignment
        max_col_count = max(max_col_count, len(headers))
        row_cursor += 1

        for row in rows:
            for col_index, value in enumerate(row, start=1):
                cell = ws.cell(row=row_cursor, column=col_index, value=value)
                cell.alignment = wrap_alignment
            row_cursor += 1
        row_cursor += 1

    for col in range(1, max_col_count + 1):
        ws.column_dimensions[get_column_letter(col)].width = 38
    for row in range(1, row_cursor + 1):
        ws.row_dimensions[row].height = 36

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export Markdown comparison reports to XLSX table files"
    )
    parser.add_argument(
        "--inputs",
        nargs="+",
        required=True,
        help="Markdown report files to export",
    )
    parser.add_argument(
        "--output-dir",
        help="Optional output directory. Defaults to each Markdown file's parent directory.",
    )
    args = parser.parse_args()

    output_dir = Path(args.output_dir).resolve() if args.output_dir else None
    generated_paths: list[Path] = []
    for item in args.inputs:
        markdown_path = Path(item).resolve()
        target_dir = output_dir or markdown_path.parent
        xlsx_path = target_dir / f"{markdown_path.stem}.xlsx"
        export_markdown_to_xlsx(markdown_path, xlsx_path)
        generated_paths.append(xlsx_path)

    for path in generated_paths:
        print(path)


if __name__ == "__main__":
    main()
