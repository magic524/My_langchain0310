from __future__ import annotations

import argparse
from pathlib import Path

from bs4 import BeautifulSoup, Tag
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt


def _iter_top_level_blocks(container: Tag) -> list[Tag]:
    blocks: list[Tag] = []
    for child in container.children:
        if not isinstance(child, Tag):
            continue
        if child.name in {"article", "header"}:
            blocks.extend(_iter_top_level_blocks(child))
            continue
        if child.name == "div" and child.get("style") == "display:contents":
            blocks.extend(_iter_top_level_blocks(child))
            continue
        blocks.append(child)
    return blocks


def _clean_text(tag: Tag) -> str:
    return tag.get_text("\n", strip=True)


def _set_run_font(run, *, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def _set_paragraph_font(paragraph) -> None:
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def _append_text_with_basic_format(paragraph, tag: Tag) -> None:
    if not list(tag.children):
        run = paragraph.add_run(tag.get_text())
        _set_run_font(run)
        return

    for child in tag.children:
        if isinstance(child, str):
            if child:
                run = paragraph.add_run(child)
                _set_run_font(run)
            continue
        if not isinstance(child, Tag):
            continue
        if child.name in {"strong", "b"}:
            run = paragraph.add_run(child.get_text())
            _set_run_font(run, bold=True)
            continue
        if child.name in {"br"}:
            paragraph.add_run("\n")
            continue
        run = paragraph.add_run(child.get_text())
        _set_run_font(run)


def _add_heading(document: Document, tag: Tag) -> None:
    level = {"h1": 1, "h2": 2, "h3": 3, "h4": 4}.get(tag.name, 2)
    paragraph = document.add_heading(level=level)
    _append_text_with_basic_format(paragraph, tag)


def _add_paragraph(document: Document, tag: Tag) -> None:
    paragraph = document.add_paragraph()
    _append_text_with_basic_format(paragraph, tag)


def _add_list(document: Document, tag: Tag) -> None:
    style_name = "List Number" if tag.name == "ol" else "List Bullet"
    for child in tag.find_all("li", recursive=False):
        paragraph = document.add_paragraph(style=style_name)
        _append_text_with_basic_format(paragraph, child)


def _add_table(document: Document, tag: Tag) -> None:
    row_tags = tag.find_all("tr")
    if not row_tags:
        return

    max_cols = max(
        len(row.find_all(["th", "td"], recursive=False))
        for row in row_tags
    )
    table = document.add_table(rows=0, cols=max_cols)
    table.style = "Table Grid"

    for row_tag in row_tags:
        cells = row_tag.find_all(["th", "td"], recursive=False)
        row = table.add_row().cells
        for index, cell_tag in enumerate(cells):
            cell = row[index]
            cell.text = _clean_text(cell_tag)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cell.paragraphs:
                _set_paragraph_font(paragraph)
                if cell_tag.name == "th":
                    for run in paragraph.runs:
                        run.bold = True


def _walk(document: Document, tag: Tag) -> None:
    if tag.name in {"h1", "h2", "h3", "h4"}:
        _add_heading(document, tag)
        return
    if tag.name == "p":
        if _clean_text(tag):
            _add_paragraph(document, tag)
        return
    if tag.name in {"ul", "ol"}:
        _add_list(document, tag)
        return
    if tag.name == "table":
        _add_table(document, tag)
        return

    for child in _iter_top_level_blocks(tag):
        _walk(document, child)


def convert_html_to_docx(input_path: Path, output_path: Path) -> None:
    soup = BeautifulSoup(input_path.read_text(encoding="utf-8"), "lxml")
    body = soup.body
    if body is None:
        raise ValueError(f"Could not find HTML body in {input_path}")

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(10.5)

    for tag in _iter_top_level_blocks(body):
        _walk(document, tag)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export simple HTML report to DOCX")
    parser.add_argument("--input", required=True, help="Input HTML file")
    parser.add_argument("--output", required=True, help="Output DOCX file")
    args = parser.parse_args()

    convert_html_to_docx(Path(args.input), Path(args.output))


if __name__ == "__main__":
    main()
